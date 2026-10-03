#!/usr/bin/env python3
"""Export sheet "dps" from mstrasset_011026.xlsx to ../data.json.

Usage:
    python3 tools/export_data.py [path/to/mstrasset_011026.xlsx]

Requires openpyxl (already available in system python3 on this machine).
"""

import datetime
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

SHEET = "dps"
DEFAULT_SRC = Path.home() / "Documents" / "mstrasset" / "Atribut Data Tanah" / "mstrasset_011026.xlsx"
OUT = Path(__file__).resolve().parent.parent / "data.json"

ID_COLS = ("kode_satker", "id_aset", "id_aset_bidang", "kd_brg", "nup", "luas", "alamat_bidang")


def norm(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    if not src.exists():
        raise SystemExit(f"Source workbook not found: {src}")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET}' not found in {src.name}: {wb.sheetnames}")
    ws = wb[SHEET]

    rows_iter = ws.iter_rows(values_only=True)
    header = [norm(h) for h in next(rows_iter)]
    while header and not header[-1]:
        header.pop()
    width = len(header)
    idx = {name: i for i, name in enumerate(header)}
    missing = [c for c in ID_COLS if c not in idx]
    if missing:
        raise SystemExit(f"Missing identity columns: {missing}")

    seen = {}
    rows = []
    for raw in rows_iter:
        values = [norm(v) for v in raw[:width]]
        values += [""] * (width - len(values))
        if not any(values):
            continue
        key = "|".join(values[idx[c]] for c in ID_COLS)
        occurrence = seen.get(key, 0) + 1
        seen[key] = occurrence
        suffix = "" if occurrence == 1 else f"#{occurrence}"
        row_id = hashlib.sha1((key + suffix).encode("utf-8")).hexdigest()[:12]
        rows.append([row_id] + values)

    payload = {
        "sheet": SHEET,
        "source": src.name,
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "columns": header,
        "rows": rows,
    }
    OUT.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(f"{len(rows)} rows x {width} columns -> {OUT} ({OUT.stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
